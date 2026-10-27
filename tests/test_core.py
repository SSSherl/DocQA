import json
from pathlib import Path
import pytest
from docqa.docbase import DocBase, chunks
from docqa.agent import answer_question
from docqa.doc2md import parse_document


@pytest.fixture
def base(tmp_path):
    return DocBase({"data_dir":str(tmp_path / "data"), "chunk_size":300, "chunk_overlap":40})


def add(base, tmp_path, name, text, category="技术"):
    file = tmp_path / name
    file.write_text(text, encoding="utf-8")
    return file, base.build(file, category)


def test_incremental_update_category_and_no_duplicates(base,tmp_path):
    file,result = add(base,tmp_path,"资料.md","# 接口\nFE_HEART_Start 用于心跳启动。")
    assert result["chunks"] == 1
    assert base.build(file)["status"] == "unchanged"
    assert len(base.query("FE_HEART_Start")["results"]) == 1
    file.write_text("# 更新\nSD5993 新型号",encoding="utf-8")
    assert base.build(file,"新版")["status"] == "indexed"
    assert not base.query("FE_HEART_Start")["results"]
    assert base.query("SD5993",category="新版")["results"]
    assert not base.query("SD5993",category="技术")["results"]


def test_chinese_traceable_and_no_evidence(base,tmp_path):
    add(base,tmp_path,"架构.md","# 三路混合检索\n采用关键词召回和融合检索。")
    result = answer_question(base,"混合检索")
    assert result["mode"] == "extractive"
    assert result["citations"][0]["name"] == "架构.md"
    assert result["citations"][0]["page"] is None
    assert answer_question(base,"ZX_UNKNOWN_987")["mode"] == "no_evidence"


def test_atomic_failure_preserves_old_index(base,tmp_path,monkeypatch):
    file,_ = add(base,tmp_path,"资料.md","旧词 OLD_TOKEN")
    file.write_text("新词 NEW_TOKEN",encoding="utf-8")
    def fail(*args):
        raise RuntimeError("模拟解析错误")
    monkeypatch.setattr("docqa.docbase.parse_document",fail)
    with pytest.raises(RuntimeError):
        base.build(file)
    assert base.query("OLD_TOKEN")["results"]
    assert not base.query("NEW_TOKEN")["results"]


def test_changed_model_config_requires_rebuild(base,tmp_path):
    add(base,tmp_path,"测试.txt","模型配置")
    base.cfg["chunk_size"] = 500
    with pytest.raises(ValueError,match="重建"):
        base.query("模型")


def test_delete_and_metrics(base,tmp_path):
    _,result = add(base,tmp_path,"检索.txt","FE_HEART_Start")
    score = base.evaluate([{"question":"FE_HEART_Start","relevant_ids":[result["id"]]}])
    assert score["recall"] == score["hit_rate"] == score["mrr"] == 1
    base.delete(result["id"])
    assert not base.documents()
    assert not base.query("FE_HEART_Start")["results"]


def test_chunks_do_not_lose_final_text():
    text = "甲"*400+"END_TOKEN"
    output = list(chunks(text,150,30))
    assert output[-1].endswith("END_TOKEN")
    with pytest.raises(ValueError):
        list(chunks(text,100,100))


def test_image_without_vision_is_not_searchable(base,tmp_path):
    from PIL import Image
    file = tmp_path / "图.png"
    Image.new("RGB",(40,40),"white").save(file)
    result = base.build(file)
    assert result["chunks"] == 0
    assert result["warnings"]
    assert list((base.root/"parsed").rglob("*.png"))


def test_office_and_pdf_provenance(tmp_path):
    from docx import Document
    from pptx import Presentation
    import pymupdf
    doc = Document()
    doc.add_heading("架构说明",level=1)
    doc.add_paragraph("正文内容")
    table = doc.add_table(rows=2,cols=2)
    table.cell(0,0).text = "接口"
    table.cell(1,0).text = "FE_HEART_Start"
    doc.save(tmp_path/"sample.docx")
    result = parse_document(tmp_path/"sample.docx",tmp_path/"word",{})
    assert any("FE_HEART_Start" in s["text"] for s in result["sections"])
    assert all(s["page"] is None for s in result["sections"])
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Architecture"
    prs.save(tmp_path/"sample.pptx")
    result = parse_document(tmp_path/"sample.pptx",tmp_path/"ppt",{})
    assert result["sections"][0]["page"] == 1
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72,72),"PDF architecture")
    page.draw_rect(pymupdf.Rect(70,90,160,140))
    pdf.save(tmp_path/"sample.pdf")
    pdf.close()
    result = parse_document(tmp_path/"sample.pdf",tmp_path/"pdf",{})
    assert "PDF architecture" in result["sections"][0]["text"]
    assert result["sections"][1]["image"]
    assert result["sections"][1]["page"] == 1


def test_model_routes_and_citation_guard(base,tmp_path,monkeypatch):
    base.cfg.update(embedding={"url":"http://127.0.0.1:9000/v1/embeddings"},reranker={"url":"http://127.0.0.1:9000/rerank"},chat={"url":"http://127.0.0.1:9000/chat"},sparse_model_path="fake-local-model")
    monkeypatch.setattr("docqa.docbase.embeddings",lambda cfg,texts:[[1.,0.] for _ in texts])
    monkeypatch.setattr(base,"sparse",lambda texts:[{"1":1.} for _ in texts])
    monkeypatch.setattr("docqa.models.post",lambda cfg,role,payload:{"results":[{"index":i,"relevance_score":1-i*.1} for i in range(len(payload["documents"]))]})
    add(base,tmp_path,"混合.md","# 混合检索\n三路召回 RRF")
    result = base.query("混合检索")
    assert set(result["routes"]) == {"BM25","dense","BGE-M3"}
    assert result["results"][0]["rerank_score"] == 1
    monkeypatch.setattr("docqa.agent.chat",lambda *args:"检索使用三路召回。[1]")
    assert answer_question(base,"混合检索")["mode"] == "model"
    monkeypatch.setattr("docqa.agent.chat",lambda *args:"无中生有。[99]")
    assert answer_question(base,"混合检索")["mode"] == "citation_rejected"


def test_model_loopback_boundary():
    from docqa.models import post
    with pytest.raises(ValueError,match="回环"):
        post({"chat":{"url":"https://remote.example/chat"}},"chat",{})
